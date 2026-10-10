#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Importar clientes desde Excel a la tabla clientes (SQLite).
La hoja ACTIVO contiene clientes; FUNCIONARIOS se omite porque no representa
clientes ni contiene NIT de empresa.
Uso: python importar_clientes_excel.py [ruta.xlsx]
"""
import sys
import os
import json
from typing import Any, Dict

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import openpyxl
from guias_coodescor.core.utils import normalizar_nit
from guias_coodescor.services.clientes_service import (
    guardar_cliente,
    obtener_cliente_por_nit,
)

EXCEL_PATH = r"D:\Users\57323\Downloads\Base de datos cliente V2.xlsx"


def _texto_excel(valor):
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor)).strip()
    return str(valor).strip()


def _metadata_existente(cliente):
    contenido = cliente.get("metadata") or "{}"
    try:
        metadata = json.loads(contenido)
    except (TypeError, ValueError):
        return None
    return metadata if isinstance(metadata, dict) else None


def importar(ruta_excel=EXCEL_PATH):
    if not os.path.exists(ruta_excel):
        print(f"[ERROR] No existe: {ruta_excel}")
        return 1

    wb = openpyxl.load_workbook(ruta_excel, read_only=True, data_only=True)
    if "ACTIVO" not in wb.sheetnames:
        wb.close()
        print("[ERROR] El libro no contiene la hoja ACTIVO.")
        return 1
    ws = wb["ACTIVO"]
    hojas_omitidas = [nombre for nombre in wb.sheetnames if nombre != "ACTIVO"]
    insertados = 0
    actualizados = 0
    sin_cambios = 0
    sin_nit = 0
    errores = 0

    try:
        # Fila 1: título; fila 2: encabezados; filas siguientes: clientes.
        for i, row in enumerate(ws.iter_rows(min_row=3, values_only=True), start=3):
            if not row or len(row) < 9:
                errores += 1
                print(f"[ERROR] Fila {i}: se esperaban al menos 9 columnas.")
                continue
            nit = normalizar_nit(_texto_excel(row[0]))
            razon = _texto_excel(row[1])
            if not nit:
                if razon:
                    sin_nit += 1
                continue
            if not razon:
                continue

            datos_excel = {
                "nit": nit,
                "razon_social": razon,
                "ciudad": _texto_excel(row[2]),
                "direccion": _texto_excel(row[3]),
                "telefono": _texto_excel(row[4]),
                "contacto": _texto_excel(row[6]),
                "email": _texto_excel(row[7]),
            }
            metadata_excel = {
                "tipo_cliente": _texto_excel(row[5]),
                "telefono_compras": _texto_excel(row[8]),
                "correo_compras": _texto_excel(row[7]),
            }
            try:
                existente = obtener_cliente_por_nit(nit)

                if existente:
                    datos_guardar: Dict[str, Any] = {
                        "nit": nit,
                        "razon_social": existente.get("razon_social") or razon,
                    }
                    metadata = _metadata_existente(existente)
                    for campo, valor in datos_excel.items():
                        if campo == "nit" or not valor:
                            continue
                        if not existente.get(campo):
                            datos_guardar[campo] = valor
                    if metadata is not None:
                        metadata_actualizada = dict(metadata)
                        for campo, valor in metadata_excel.items():
                            if valor and not metadata_actualizada.get(campo):
                                metadata_actualizada[campo] = valor
                        if metadata_actualizada != metadata:
                            datos_guardar["metadata"] = metadata_actualizada
                    if len(datos_guardar) == 2:
                        sin_cambios += 1
                        continue
                else:
                    datos_guardar: Dict[str, Any] = {"nit": nit, "razon_social": razon}
                    datos_guardar.update({
                        campo: valor for campo, valor in datos_excel.items()
                        if campo != "nit" and valor
                    })
                    datos_guardar["metadata"] = {
                        campo: valor for campo, valor in metadata_excel.items() if valor
                    }

                ok, mensaje, _ = guardar_cliente(datos_guardar, origen="manual")
                if not ok:
                    errores += 1
                    print(f"[ERROR] Fila {i}, NIT {nit}: {mensaje}")
                elif existente:
                    actualizados += 1
                else:
                    insertados += 1
            except Exception as exc:
                errores += 1
                print(f"[ERROR] Fila {i}, NIT {nit}: {exc}")
    finally:
        wb.close()

    if hojas_omitidas:
        print(f"[INFO] Hojas omitidas (no son clientes): {', '.join(hojas_omitidas)}")
    print(
        f"[OK] Insertados: {insertados}, actualizados: {actualizados}, "
        f"sin cambios: {sin_cambios}, filas sin NIT: {sin_nit}, errores: {errores}"
    )
    return 1 if errores else 0


if __name__ == "__main__":
    sys.exit(importar(sys.argv[1] if len(sys.argv) > 1 else EXCEL_PATH))