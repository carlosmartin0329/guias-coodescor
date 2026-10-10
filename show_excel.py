import openpyxl
wb = openpyxl.load_workbook(r'D:\Users\57323\Downloads\Base de datos cliente V2.xlsx', read_only=True)
ws = wb.active
print('Sheet:', ws.title)
print('Max row:', ws.max_row, 'Max col:', ws.max_column)

# Headers
headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
print('Headers:', headers)
print()

# First 10 rows
for i, row in enumerate(ws.iter_rows(min_row=2, max_row=11, values_only=True)):
    print('Row {}: {}'.format(i+1, row))

wb.close()